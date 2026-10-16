"""Build the editable Group 6 AlignAgent strawman deck with python-pptx.

Run: python3 presentation/build_ghost_deck.py
Dependencies: python-pptx (or the local presentation/.dependencies directory).
All generated files remain beside this script. No project files are modified.
"""
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / '.dependencies'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement

OUT = ROOT / 'Group6_AlignAgent_Ghost_Deck.pptx'
NAVY = '142C43'
TEAL = '087F8C'
GRAY = '526171'
LIGHT = 'EAF0F4'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
prs.core_properties.title = 'Group 6 | AlignAgent Communication Alignment AI Agent'
prs.core_properties.subject = 'Ghost / strawman deck; proposed design and pending evidence'
prs.core_properties.author = 'Group 6'


def text(slide, content, x, y, w, h, size=21, color=NAVY, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for i, line in enumerate(content.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Arial'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(10)
    return shape


def base(title, subtitle='Proposed approach • Evidence and implementation pending'):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, title, .65, .48, 12.0, .95, 32, bold=True)
    text(s, subtitle, .65, 1.48, 12, .45, 17, GRAY)
    text(s, 'GROUP 6  |  ALIGNAGENT  |  STRAWMAN', .65, 7.03, 10, .25, 10, GRAY)
    text(s, f'{len(prs.slides):02d} / 19', 11.55, 7.03, 1.1, .25, 11, GRAY)
    return s


def bullets(s, items, x=.75, y=2.18, w=11.8, size=23, step=.86):
    for i, item in enumerate(items):
        assert item.endswith(('.', '?', '!')), item
        shape = text(s, item, x, y + i * step, w, step-.12, size)
        p = shape.text_frame.paragraphs[0]
        props = p._p.get_or_add_pPr()
        props.set('marL', '228600')
        props.set('indent', '-228600')
        char = OxmlElement('a:buChar')
        char.set('char', '•')
        props.append(char)


def table(s, caption, headers, rows, widths=None, y=2.6, height=3.6, size=18):
    text(s, caption, .7, 2.05, 12, .45, 18, TEAL, True)
    tbl = s.shapes.add_table(len(rows)+1, len(headers), Inches(.7), Inches(y), Inches(11.9), Inches(height)).table
    if widths:
        for col, width in zip(tbl.columns, widths):
            col.width = Inches(width)
    for r, row in enumerate([headers]+rows):
        for c, value in enumerate(row):
            cell = tbl.cell(r,c)
            cell.text = value
            cell.margin_left = cell.margin_right = Inches(.13)
            cell.margin_top = cell.margin_bottom = Inches(.09)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(NAVY if r == 0 else ('F1F5F7' if r%2 else 'FFFFFF'))
            for p in cell.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(size)
                p.font.bold = r == 0
                p.font.color.rgb = RGBColor.from_string('FFFFFF' if r == 0 else NAVY)
    return tbl


def note(s, value):
    s.notes_slide.notes_text_frame.text = value


def content(title, items, subtitle=None):
    s = base(title, subtitle) if subtitle else base(title)
    bullets(s, items)
    note(s, 'Strawman proposal, not a statement of implemented capability. Replace placeholders with approved project evidence before final submission.')
    return s


s = base('AlignAgent', 'Communication Alignment AI Agent')
text(s, 'Group 6', .75, 2.55, 11.8, .6, 28, TEAL, True)
text(s, 'Ghost / strawman project deck', .75, 3.4, 11.8, .6, 27)
text(s, '[Presenter names]  •  [Course / sponsor]  •  [Presentation date]', .75, 4.5, 11.8, .8, 20, GRAY)
text(s, 'Working proposal. Research, performance results, and business impact remain unverified.', .75, 5.65, 11.5, .8, 20, GRAY)

content('Executive Summary', [
    'Proposed purpose: help users check whether messages reflect an agreed goal and audience.',
    'Personal mode supports drafting; organizational mode checks approved priorities and policies.',
    'Outputs would explain possible gaps, cite supporting context, and suggest revisions.',
    'Decision sought: agree on scope, data access, and an evaluation plan before a pilot.',
    '[Pending: validated need, prototype status, evaluation findings, and sponsor feedback].',
])
content('Problem Understanding', [
    'Working hypothesis: unclear goals and inconsistent language can create communication gaps.',
    'Candidate gaps: missing requests, ambiguous ownership, conflicting priorities, or unsuitable tone.',
    'Alignment means consistency with an explicit goal, intended audience, and approved context.',
    'Validate through interviews and consented examples; distinguish preference from material error.',
    '[Pending: problem statement, observed examples, frequency, and affected stakeholders].',
])
content('Business Need', [
    'Proposed need: reduce avoidable clarification and help teams communicate decisions clearly.',
    'Baseline measures: drafting time, clarification rounds, and reviewer-rated message quality.',
    'Sponsor and decision owner: [Pending]; target workflow and pilot population: [Pending].',
    'Adoption depends on useful feedback, source traceability, and clear user control.',
    'Business case remains a hypothesis until benefits and operating costs are measured.',
])
content('Proposed AlignAgent Solution', [
    'Capture the draft, intended audience, communication goal, and user-approved context.',
    'Identify potential gaps against a transparent rubric and show supporting passages.',
    'Suggest revisions while preserving the author’s meaning and final editorial control.',
    'Use retrieval-augmented generation (RAG): retrieve relevant context before generating feedback.',
    'Use Structured Query Language (SQL) for authorized structured records and evaluation logs.',
])
s = base('User Personas and Use Cases')
table(s, 'Table 1. Candidate personas and workflows; validation pending.',
      ['Persona', 'Candidate use case', 'Expected output'], [
          ['Individual contributor', 'Draft a request or status update.', 'Clear request and suggested revision.'],
          ['Team lead', 'Review a decision announcement.', 'Goal coverage and unresolved questions.'],
          ['Communications owner', 'Check a policy announcement.', 'Cited inconsistencies with approved text.'],
          ['Administrator / reviewer', 'Review access and quality.', 'Authorized audit and evaluation records.'],
      ], widths=[2.6,4.6,4.7])
content('Methodology', [
    'Discovery: interview intended users and document current communication workflows.',
    'Research: review relevant literature and collect permitted, de-identified examples.',
    'Design: define the alignment rubric, data boundaries, and review responsibilities.',
    'Prototype: implement a narrow workflow and compare against simpler baselines.',
    'Evaluate: test held-out examples, review failures, and revise before a pilot.',
])
s = base('System Architecture', 'Conceptual design • Components and integrations are not implementation claims')
text(s, 'Figure 1. Proposed request and evidence flow.', .7, 2.05, 12, .45, 18, TEAL, True)
layers = [
    ('User input', 'Draft + goal + audience + consent'),
    ('Access and retrieval', 'Permission checks → RAG context / authorized SQL records'),
    ('Alignment analysis', 'Rubric checks → cited feedback → optional revision'),
    ('Human review', 'Author accepts, edits, or rejects suggestions'),
]
for i, (label, detail) in enumerate(layers):
    yy = 2.75 + i*.87
    text(s, label, .9, yy, 3.1, .58, 22, TEAL, True)
    text(s, detail, 4.1, yy, 8.0, .65, 21)
text(s, 'Cross-cutting controls: data minimization, audit records, and source versioning.', .9, 6.35, 11.7, .4, 18, GRAY)
note(s, 'Figure 1 is an editable conceptual architecture. Provider, storage, deployment, and integrations are pending design decisions.')
content('Data and Research Sources', [
    'Candidate project inputs: approved policies, communication goals, and consented message examples.',
    'Candidate research: peer-reviewed work on communication quality and grounded generation.',
    'Record source owner, permission, date, version, retention rule, and relevance for every input.',
    'Separate synthetic demonstration inputs from real research and evaluation data.',
    '[Pending: source inventory, access approvals, sample size, and completed literature review].',
])
content('Analytics and AI Techniques', [
    'Rubric checks: goal coverage, clarity, actionable requests, and consistency with context.',
    'Semantic retrieval: rank relevant approved passages for RAG; assess retrieval relevance.',
    'Language-model analysis: generate explanations and revisions with traceable citations.',
    'SQL analytics: summarize authorized usage and evaluation records without exposing drafts.',
    'Compare against manual review and rubric-only prompting; model selection remains pending.',
])
content('Personal-Use Agent', [
    'Scenario: an individual prepares a project update for a specified audience.',
    'Inputs: draft, communication goal, optional user-provided context, and tone preference.',
    'Proposed output: missing information, explained suggestions, and an editable revised draft.',
    'The user chooses what to submit and retains control over changes and external sharing.',
    '[Pending: prototype walkthrough, consent controls, and usability observations].',
])
content('Organizational Agent', [
    'Scenario: a communications owner reviews an announcement against approved priorities.',
    'Inputs: draft, audience, current approved sources, and role-based access permissions.',
    'Proposed output: source-linked discrepancies and questions for the responsible reviewer.',
    'Use approved source versions and prohibit retrieval outside the requester’s permissions.',
    '[Pending: sponsor requirements, governance owner, integration scope, and pilot approval].',
])
s = base('Evaluation Framework / Research Evidence Matrix', 'Planned evaluation • No scores or acceptance thresholds have been established')
table(s, 'Table 2. Evidence needed to assess the proposal.',
      ['Question', 'Method / measure', 'Evidence / status'], [
          ['Is feedback correct?', 'Expert rubric; citation support rate.', '[Pending: labeled test set and scores].'],
          ['Does it help authors?', 'Task time; blind review of revisions.', '[Pending: consented user study].'],
          ['Is retrieval relevant?', 'Relevant passages retrieved / expected.', '[Pending: reference passages].'],
          ['Are controls effective?', 'Access denial and prompt-injection tests.', '[Pending: test protocol and findings].'],
          ['Is it economical?', 'Cost per task; response latency.', '[Pending: workload and measurements].'],
      ], widths=[2.7,4.35,4.85], height=3.9, size=17)
note(s, 'Predefine sample size, annotation instructions, thresholds, and adjudication. Use held-out examples; document disagreement and uncertainty. Do not treat model self-grading as independent research evidence.')
s = base('Results and Demo Plan', 'Results pending • Empty chart placeholder; no observed or simulated values')
bullets(s, [
    'Demo: enter a synthetic draft and an explicit communication goal.',
    'Show retrieved context, explained feedback, and an author-reviewed revision.',
    '[Pending: completed runs, failure examples, and measured results].',
], w=5.6, size=21, step=1.0)
text(s, 'Figure 2. Planned quality comparison.', 7.0, 2.1, 5.4, .5, 18, TEAL, True)
for xa,ya,xb,yb in [(7.75,5.65,12.1,5.65),(7.75,5.65,7.75,3.05)]:
    line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(xa), Inches(ya), Inches(xb), Inches(yb))
    line.line.color.rgb = RGBColor.from_string(GRAY)
    line.line.width = Pt(1.5)
text(s, 'Reviewer-rated alignment\n(score; scale pending)', 7.8, 2.65, 4.4, .7, 17, GRAY)
text(s, '[No data available]', 8.3, 4.25, 3.5, .5, 21, GRAY)
text(s, 'Original', 7.9, 5.8, 1.8, .4, 17)
text(s, 'Agent-assisted', 10.0, 5.8, 2.5, .4, 17)
text(s, 'Message condition', 8.65, 6.35, 3.5, .4, 18, GRAY)
content('Business Value and Recommendations', [
    'Hypothesized value: clearer requests, fewer revision cycles, and less reviewer effort.',
    'Estimate value as task volume × verified time saved × labor cost, less operating costs.',
    'Recommendation: start with one approved workflow and an accountable business owner.',
    'Expand only if quality, privacy controls, and measured economics meet agreed criteria.',
    '[Pending: baseline values, cost assumptions, evaluation outcomes, and sponsor decision].',
])
content('Risks, Assumptions, and Limitations', [
    'Risk: unsupported feedback; mitigate with source citations and explicit uncertainty.',
    'Risk: sensitive content exposure; minimize collection and enforce access boundaries.',
    'Risk: biased tone judgments; test diverse audiences and allow authors to reject advice.',
    'Assumption: approved sources reflect current goals; assign owners and review dates.',
    'Limitations: no validated dataset, completed evaluation, or proven business impact yet.',
])
content('Potential Next Steps', [
    'Agree on the first use case and define who can approve scope and source access.',
    'Complete interviews, source inventory, and a literature review with traceable citations.',
    'Define the rubric and held-out evaluation set before implementing the prototype.',
    'Run a controlled comparison and document errors alongside any measured benefits.',
    '[Pending: milestone dates, task owners, pilot criteria, and go / no-go decision].',
])
s = base('Appendix', 'Working definitions and completion checklist')
table(s, 'Table 3. Terms and evidence to complete.', ['Item', 'Definition / required addition'], [
    ['Alignment', 'Consistency with the stated goal, audience, and approved context.'],
    ['RAG', 'Retrieval-augmented generation; retrieve context before generation.'],
    ['SQL', 'Structured Query Language; query permitted structured records.'],
    ['Evidence package', '[Pending: rubric, source inventory, study protocol, and test set].'],
    ['Demo package', '[Pending: synthetic inputs, walkthrough, and recorded failure cases].'],
], widths=[2.6,9.3], height=3.9, size=18)
s = content('Bibliography', [
    '[B1 Pending: communication alignment research; author, year, title, venue, and DOI / URL].',
    '[B2 Pending: original RAG research; verified publication details and DOI / URL].',
    '[B3 Pending: selected model and retrieval documentation; version, URL, and access date].',
    '[B4 Pending: applicable data governance sources; issuer, date, title, and URL].',
    '[B5 Pending: project interviews and approved documents; owner, date, and access basis].',
], 'Research placeholder register • No completed literature review is claimed')
note(s, 'No external research was used to claim project outcomes. Replace B1–B5 with verified references and cite them on the relevant slides. Confirm all source metadata before final submission.')


def validate():
    assert len(prs.slides) == 19
    for index, slide in enumerate(prs.slides, 1):
        assert any(f'{index:02d} / 19' in shape.text for shape in slide.shapes if shape.has_text_frame)
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width
            assert shape.top + shape.height <= prs.slide_height
    prs.save(OUT)
    reopened = Presentation(OUT)
    assert len(reopened.slides) == 19 and OUT.stat().st_size > 0
    print(f'Created and verified: {OUT} ({OUT.stat().st_size:,} bytes; 19 slides).')


if __name__ == '__main__':
    validate()
